"""Builds tiny real PDF and DOCX files in memory for extraction tests."""

from io import BytesIO

from docx import Document
from pypdf import PdfWriter


def text_pdf(text: str) -> bytes:
    """A one-page PDF whose content stream draws `text` in Helvetica."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_at,
    )
    return bytes(out)


def blank_pdf() -> bytes:
    """A valid PDF with no text, like a scanned image resume."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    return _write_pdf(writer)


def encrypted_pdf() -> bytes:
    writer = PdfWriter(clone_from=BytesIO(text_pdf("Secret resume")))
    writer.encrypt(user_password="secret", algorithm="RC4-128")
    return _write_pdf(writer)


def docx(paragraphs: list[str], table_cells: list[str] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_cells:
        table = document.add_table(rows=1, cols=len(table_cells))
        for cell, text in zip(table.rows[0].cells, table_cells):
            cell.text = text
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _write_pdf(writer: PdfWriter) -> bytes:
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
