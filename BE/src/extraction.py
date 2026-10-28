"""Turns an uploaded PDF or DOCX resume into plain text."""

import logging
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from src.errors import AppError, ErrorCode

logger = logging.getLogger(__name__)

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def extract_resume_text(
    filename: str | None, content_type: str | None, data: bytes
) -> str:
    """Return the resume's text, or raise AppError explaining why it can't be read."""
    name = (filename or "").lower()
    if content_type == PDF_CONTENT_TYPE or name.endswith(".pdf"):
        read = _read_pdf
    elif content_type == DOCX_CONTENT_TYPE or name.endswith(".docx"):
        read = _read_docx
    else:
        raise AppError(
            ErrorCode.UNSUPPORTED_FILE_TYPE,
            415,
            "Only PDF and DOCX resumes are supported. "
            "Save your resume as one of those and upload it again.",
        )

    if not data:
        raise _unreadable(
            "The resume file is empty. Choose the file again and upload it."
        )

    try:
        text = read(data)
    except Exception as exc:  # parsers raise many types on malformed input
        logger.info("Could not parse resume %r: %r", filename, exc)
        raise _unreadable(
            "The resume could not be opened. It may be damaged or password "
            "protected; export it again as an unlocked PDF or DOCX."
        ) from exc

    if not text.strip():
        raise _unreadable(
            "The resume has no selectable text, which usually means it is a "
            "scanned image. Upload a text-based PDF or DOCX."
        )
    return text


def _read_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(data: bytes) -> str:
    document = Document(BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def _unreadable(detail: str) -> AppError:
    return AppError(ErrorCode.RESUME_UNREADABLE, 422, detail)
